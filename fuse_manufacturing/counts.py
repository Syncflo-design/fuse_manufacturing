"""Loading a counted Intacct cycle count sheet back into Intacct.

The count belongs to Intacct from start to finish. Someone creates it there, starts it,
prints or downloads the worksheet, and reconciles it there when the counting is done.
Fuse does one thing in the middle: it takes the worksheet back once the counters have
filled it in, and writes each counted quantity onto its line in Intacct. That saves someone
from keying hundreds of lines in by hand, twice.

Twice because the process is count, look at the variances, recount the lines that are out,
and load again. The second sheet only has figures on the recounted lines. A line left
blank is NOT sent, so the recount leaves the first count alone everywhere else. Blank means
"not counted on this sheet". A zero means zero, and is sent.

What the API allows, from Intacct's own documentation, not assumed:
  * ICCYCLECOUNTENTRY can be UPDATED with QUANTITYCOUNTED, QUANTITYDAMAGED,
    ADJUSTMENTREASON and COUNTEDBYID, keyed on RECORDNO.
  * Not while the count is Not Started. It has to be started in Intacct first.
  * ICCYCLECOUNT itself cannot be updated. There is no API call to start, finish or
    reconcile a count, which is why those steps stay in Intacct.

Every value is sent as an absolute figure, never a difference. Loading the same sheet twice
leaves Intacct exactly as loading it once did, and that is what makes a retry safe.
"""

import csv
import io
import os
import xml.etree.ElementTree as ET

import frappe
from frappe.utils import cstr
from fuse_core import gateway

# The worksheet's labels, exactly as Intacct writes them. The header block is a label in the
# first column and its value in the second, then a row of column headings, then the lines.
HEADER_LABELS = {
	"count id:": "count_id",
	"count description:": "description",
	"warehouse:": "warehouse",
}

COLUMNS = {
	"item id": "item_id",
	"item name": "item_name",
	"bin": "bin",
	"row": "row",
	"aisle": "aisle",
	"zone": "zone",
	"unit": "unit",
	"qty counted": "qty_counted",
	"qty damaged": "qty_damaged",
}

REQUIRED_COLUMNS = ("item_id", "unit", "qty_counted")

# Updates per request. A size chosen to keep one request small, not a limit Intacct
# publishes. Each request is atomic, so a failure never half-applies one.
BATCH = 100

# Problems are listed in full up to this many. Past that, one more line says how many were
# left out, rather than a message nobody can scroll through.
SHOW_PROBLEMS = 25


# ──────────────────────────────────────────────────────────────────────────────
# Reading the sheet
# ──────────────────────────────────────────────────────────────────────────────


def read_file(file_url):
	"""Every row of the first sheet, as lists of raw cell values.

	Accepts the worksheet as Intacct produces it (.xlsb), as Excel saves it (.xlsx), and CSV.
	The legacy .xls format is refused by name, because a library that half-reads it would be
	worse than being told to save it again.
	"""
	file_doc = frappe.get_doc("File", {"file_url": file_url})
	content = file_doc.get_content()
	if isinstance(content, str):
		content = content.encode("utf-8")

	extension = os.path.splitext(file_doc.file_name or file_url)[1].lower()

	if extension == ".xlsb":
		from pyxlsb import open_workbook

		with open_workbook(io.BytesIO(content)) as workbook:
			with workbook.get_sheet(1) as sheet:
				return [[cell.v for cell in row] for row in sheet.rows()]

	if extension == ".xlsx":
		from openpyxl import load_workbook

		workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
		try:
			return [list(row) for row in workbook.worksheets[0].iter_rows(values_only=True)]
		finally:
			workbook.close()

	if extension == ".csv":
		text = content.decode("utf-8-sig")
		# Excel on a South African machine saves CSV with semicolons. Detect it, do not
		# assume a comma.
		try:
			dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
		except csv.Error:
			dialect = csv.excel
		return [row for row in csv.reader(io.StringIO(text), dialect)]

	frappe.throw(
		f"Fuse cannot read a {extension or 'file with no'} extension. Attach the worksheet as "
		"Intacct downloads it (.xlsb), or save it from Excel as .xlsx or .csv.",
		title="File type not supported",
	)


def _text(value):
	"""A cell as text. Excel holds 11003150 as the number 11003150.0, not the item ID."""
	if value is None:
		return ""
	if isinstance(value, float) and value.is_integer():
		return str(int(value))
	return cstr(value).strip()


def _quantity(value):
	"""A counted figure, or None where the cell is blank. Raises ValueError on anything else.

	A CSV saved on a machine set to a decimal comma holds 107,9 for 107.9. Taken as a decimal
	comma only when there is no point in it as well, so 1,250.5 is never misread.
	"""
	if value is None:
		return None
	if isinstance(value, (int, float)):
		return float(value)

	text = cstr(value).strip().replace(" ", "")
	if not text:
		return None
	if "," in text and "." not in text:
		text = text.replace(",", ".")
	else:
		text = text.replace(",", "")
	return float(text)


def _at(values, columns, name, default):
	"""The value under one heading, or `default` where the column or the cell is missing."""
	position = columns.get(name)
	if position is None or position >= len(values):
		return default
	return values[position]


def parse_sheet(rows):
	"""The worksheet's header and lines, plus a list of what is wrong with it.

	Columns are found by their heading, not their position. A counter who hides or moves a
	column has not changed what the column means.
	"""
	header = {}
	columns = None
	lines = []
	problems = []

	for index, raw in enumerate(rows, start=1):
		cells = [_text(value) for value in raw]
		if not any(cells):
			continue

		if columns is None:
			label = cells[0].lower()
			if label in HEADER_LABELS:
				header[HEADER_LABELS[label]] = cells[1] if len(cells) > 1 else ""
				continue
			if label == "item id":
				columns = {
					COLUMNS[heading.lower()]: position
					for position, heading in enumerate(cells)
					if heading.lower() in COLUMNS
				}
				missing = [name for name in REQUIRED_COLUMNS if name not in columns]
				if missing:
					problems.append(
						"The sheet has no "
						+ ", ".join(f"'{name.replace('_', ' ').title()}'" for name in missing)
						+ " column. Use the worksheet exactly as Intacct downloads it."
					)
					return header, [], problems
			continue

		line = {
			"sheet_row": index,
			"item_id": _at(cells, columns, "item_id", ""),
			"item_name": _at(cells, columns, "item_name", ""),
			"bin": _at(cells, columns, "bin", ""),
			"row": _at(cells, columns, "row", ""),
			"aisle": _at(cells, columns, "aisle", ""),
			"zone": _at(cells, columns, "zone", ""),
			"unit": _at(cells, columns, "unit", ""),
		}

		for name in ("qty_counted", "qty_damaged"):
			try:
				line[name] = _quantity(_at(raw, columns, name, None))
			except ValueError:
				problems.append(f"Row {index}: '{_at(cells, columns, name, '')}' is not a quantity.")
				line[name] = None

		if not line["item_id"]:
			if line["qty_counted"] is not None or line["qty_damaged"] is not None:
				problems.append(f"Row {index}: there is a quantity but no Item ID.")
			continue

		for name in ("qty_counted", "qty_damaged"):
			if line[name] is not None and line[name] < 0:
				problems.append(f"Row {index} ({line['item_id']}): a count cannot be negative.")

		if line["qty_damaged"] is not None and line["qty_counted"] is None:
			problems.append(
				f"Row {index} ({line['item_id']}): damaged is filled in but counted is blank. "
				"Enter the counted quantity as well, even if it is 0."
			)

		lines.append(line)

	if columns is None:
		problems.append(
			"This does not look like an Intacct count worksheet: there is no 'Item ID' heading. "
			"Download the worksheet from the count in Intacct and fill that in."
		)
	elif not header.get("count_id"):
		problems.append(
			"The sheet has no Count ID in its header, so Fuse cannot tell which Intacct count "
			"it belongs to. Use the worksheet exactly as Intacct downloads it."
		)

	return header, lines, problems


# ──────────────────────────────────────────────────────────────────────────────
# Intacct
# ──────────────────────────────────────────────────────────────────────────────


def _equal(field, value):
	element = ET.Element("equalto")
	ET.SubElement(element, "field").text = field
	ET.SubElement(element, "value").text = value
	return ET.tostring(element, encoding="unicode")


def read_count(count_id, company):
	"""The count's header from Intacct, or None if Intacct has no count by that ID."""
	rows = gateway.query(
		"ICCYCLECOUNT",
		["RECORDNO", "CYCLECOUNTID", "CYCLECOUNTDESC", "WAREHOUSEID", "COUNTSTATE", "SHOWQTYONHAND"],
		filter_xml=_equal("CYCLECOUNTID", count_id),
		company=company,
	)
	if not rows:
		return None
	row = rows[0]
	return {
		"record_no": gateway.val(row, "RECORDNO"),
		"count_id": gateway.val(row, "CYCLECOUNTID"),
		"description": gateway.val(row, "CYCLECOUNTDESC") or "",
		"warehouse": gateway.val(row, "WAREHOUSEID") or "",
		"state": gateway.val(row, "COUNTSTATE") or "",
		"show_on_hand": gateway.flag(row, "SHOWQTYONHAND"),
	}


def read_count_lines(record_no, company):
	"""Every line on the count, as Intacct holds it."""
	rows = gateway.query(
		"ICCYCLECOUNTENTRY",
		[
			"RECORDNO", "ITEMID", "ITEMNAME", "BINID", "ICROWID", "AISLEID", "ZONEID",
			"ITEMUNIT", "QUANTITYONHAND", "LOTNO", "SERIALNO",
		],
		filter_xml=_equal("CYCLECOUNTKEY", record_no),
		company=company,
	)
	return [
		{
			"record_no": gateway.val(row, "RECORDNO"),
			"item_id": gateway.val(row, "ITEMID") or "",
			"item_name": gateway.val(row, "ITEMNAME") or "",
			"bin": gateway.val(row, "BINID") or "",
			"row": gateway.val(row, "ICROWID") or "",
			"aisle": gateway.val(row, "AISLEID") or "",
			"zone": gateway.val(row, "ZONEID") or "",
			"unit": gateway.val(row, "ITEMUNIT") or "",
			"on_hand": gateway.number(row, "QUANTITYONHAND", 0),
			"tracked": bool(gateway.val(row, "LOTNO") or gateway.val(row, "SERIALNO")),
		}
		for row in rows
	]


def _key(line):
	"""Where a line sits: the item, and the bin location if the warehouse uses them."""
	return (line["item_id"], line["bin"], line["row"], line["aisle"], line["zone"])


def _where(line):
	location = " / ".join(part for part in (line["bin"], line["row"], line["aisle"], line["zone"]) if part)
	return f"{line['item_id']} in {location}" if location else line["item_id"]


def match(sheet_lines, count_lines):
	"""Pair each sheet line with the Intacct line it was printed from.

	Refuses rather than guesses:
	  * a sheet line with no Intacct line — the sheet is for a different count, or a line
	    was typed in. Either way nothing on the count can take it.
	  * two Intacct lines at the same place. That happens on lot- or serial-tracked items,
	    one line per lot, and the worksheet carries no lot column to tell them apart.
	  * a sheet line appearing twice, which would send two figures for one line.
	  * a unit that differs from Intacct's. Counting in the wrong unit is a wrong count.
	"""
	by_key = {}
	for line in count_lines:
		by_key.setdefault(_key(line), []).append(line)

	matched = []
	problems = []
	seen = {}

	for line in sheet_lines:
		key = _key(line)
		candidates = by_key.get(key, [])

		if not candidates:
			problems.append(
				f"Row {line['sheet_row']}: {_where(line)} is not on this count in Intacct."
			)
			continue
		if len(candidates) > 1:
			problems.append(
				f"Row {line['sheet_row']}: {_where(line)} is on the count {len(candidates)} times, "
				"once per lot or serial number, and the worksheet does not say which. Enter "
				"these lines in Intacct."
			)
			continue
		if key in seen:
			problems.append(
				f"Row {line['sheet_row']}: {_where(line)} is already on row {seen[key]}. "
				"Each line may appear once."
			)
			continue
		seen[key] = line["sheet_row"]

		intacct = candidates[0]
		if line["unit"] and line["unit"] != intacct["unit"]:
			problems.append(
				f"Row {line['sheet_row']}: {_where(line)} is counted in {intacct['unit']} in "
				f"Intacct, not {line['unit']}."
			)
			continue

		matched.append(dict(line, intacct=intacct))

	return matched, problems


def raise_problems(problems, title):
	shown = problems[:SHOW_PROBLEMS]
	if len(problems) > len(shown):
		shown.append(f"… and {len(problems) - len(shown)} more.")
	frappe.throw("<br>".join(frappe.utils.escape_html(problem) for problem in shown), title=title)


def _format(quantity):
	"""A plain decimal. Never exponent notation, never a trailing .0 Intacct might query."""
	return f"{quantity:.6f}".rstrip("0").rstrip(".")


def send(lines, company, reference, stamp):
	"""Write each counted quantity onto its Intacct line. Returns how many were sent.

	`lines` are the document's rows that carry a count. `stamp` makes each attempt's control
	IDs distinct: the values are absolute, so a replay changes nothing, and a retry after one
	batch failed must not be refused because an earlier batch already used its ID.
	"""
	functions = []
	for line in lines:
		# execute_many takes <function> elements and stamps each with its control ID.
		function = ET.Element("function")
		entry = ET.SubElement(ET.SubElement(function, "update"), "ICCYCLECOUNTENTRY")
		ET.SubElement(entry, "RECORDNO").text = cstr(line.intacct_line)
		ET.SubElement(entry, "QUANTITYCOUNTED").text = _format(line.qty_counted)
		if line.damaged_entered:
			ET.SubElement(entry, "QUANTITYDAMAGED").text = _format(line.qty_damaged)
		functions.append(function)

	for start in range(0, len(functions), BATCH):
		gateway.execute_many(
			functions[start : start + BATCH],
			company=company,
			reference=reference,
			purpose=f"count-{stamp}-{start // BATCH + 1}",
			atomic=True,
		)

	return len(functions)
